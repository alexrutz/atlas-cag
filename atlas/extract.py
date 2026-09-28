"""Plain-text extraction for uploaded documents."""

import io
import re
from html.parser import HTMLParser
from pathlib import PurePath

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".xml", ".yaml", ".yml",
    ".log", ".ini", ".toml", ".cfg", ".sql", ".py", ".js", ".ts", ".java", ".c", ".h", ".cpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".tex",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".html", ".htm"}


class ExtractionError(ValueError):
    pass


class _HTMLText(HTMLParser):
    _BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}
    _SKIP = {"script", "style", "noscript", "template"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip += 1
        elif tag in self._BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP:
            self._skip = max(0, self._skip - 1)
        elif tag in self._BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.out.append(data)


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    if b"\x00" in data[:4096]:
        raise ExtractionError("file looks binary")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Page {i}]\n{text}")
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    blocks = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n\n".join(blocks)


def _html(data: bytes) -> str:
    p = _HTMLText()
    p.feed(_decode(data))
    return "".join(p.out)


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> str:
    ext = PurePath(filename).suffix.lower()
    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext in (".html", ".htm"):
            text = _html(data)
        elif ext in TEXT_EXTENSIONS or not ext:
            text = _decode(data)
        else:
            raise ExtractionError(f"unsupported file type '{ext}'")
    except ExtractionError:
        raise
    except Exception as e:  # parser libraries raise a zoo of exception types
        raise ExtractionError(f"could not parse {ext or 'file'}: {e}") from e
    text = normalize(text)
    if not text:
        raise ExtractionError("no extractable text (scanned PDFs need OCR before ingestion)")
    return text
